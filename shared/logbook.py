from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import warnings

from .experiment_meta import (
    AAT_COUNTERBALANCE_A,
    get_logbook_path,
    assign_aat_counterbalance,
    assign_iat_version,
    parse_participant_last4,
    smell_condition_order,
    validate_condition,
    validate_smell_order_group,
    validate_script_type,
    validate_task,
    validate_task_order_group,
    SMELL_ORDER_CLEAN_THEN_ISO,
    SMELL_ORDER_ISO_THEN_CLEAN,
    TASK_ORDER_AAT_BLOCK_FIRST,
    TASK_ORDER_IAT_BLOCK_FIRST,
)


SHEET_PARTICIPANTS = "Participants"
SHEET_RUNS = "Runs"
SHEET_IAT_BALANCE = "IAT Balance"
SHEET_AAT_BALANCE = "AAT Balance"
SHEET_COUNTERBALANCING_KEY = "Counterbalancing Key"
SHEET_PARTICIPANT_PLAN = "Participant Plan"
REQUIRED_SHEETS = (
    SHEET_PARTICIPANTS,
    SHEET_RUNS,
    SHEET_IAT_BALANCE,
    SHEET_AAT_BALANCE,
    SHEET_COUNTERBALANCING_KEY,
    SHEET_PARTICIPANT_PLAN,
)

PARTICIPANTS_HEADERS = [
    "participant_id",
    "participant_display_id",
    "participant_name_or_alias",
    "session_1_researcher",
    "session_2_researcher",
    "session_3_researcher",
    "phone_last4",
    "session_1_date",
    "session_2_date",
    "session_3_date",
    "session_1_condition",
    "session_2_condition",
    "session_3_condition",
    "notes",
    "smell_order_group",
    "task_order_group",
    "counterbalance_code",
]

RUNS_HEADERS = [
    "run_id",
    "participant_id",
    "participant_alias",
    "task",
    "script_type",
    "task_order_group",
    "run_timestamp",
    "session_number",
    "researcher",
    "condition",
    "iat_version",
    "iat_order_condition",
    "iat_side_condition",
    "aat_counterbalance",
    "output_filename",
    "output_path",
    "status",
    "notes",
    "smell_order_group",
    "counterbalance_code",
]

LEGACY_PARTICIPANTS_HEADERS = PARTICIPANTS_HEADERS[:-1]
LEGACY_RUNS_HEADERS = RUNS_HEADERS[:-1]
PARTICIPANT_PLAN_HEADERS = [
    "participant_id",
    "participant_alias",
    "counterbalance_code",
    "smell_order_group",
    "odor_1",
    "odor_2",
    "odor_3",
    "task_order_group",
    "iat_version",
    "iat_order_condition",
    "iat_side_condition",
    "aat_counterbalance",
    "aat_block_order",
    "run_count",
    "latest_status",
    "notes",
]
COUNTERBALANCING_KEY_HEADERS = [
    "counterbalance_code",
    "task_order_group",
    "smell_order_group",
    "odor_1",
    "odor_2",
    "odor_3",
    "iat_version",
    "iat_order_condition",
    "iat_side_condition",
    "aat_counterbalance",
    "aat_block_order",
]

VALID_RUN_STATUSES = {"completed", "interrupted", "invalid"}
BALANCE_HEADERS = ["metric", "value", "count"]
ASSIGNABLE_SMELL_ORDER_GROUPS = [
    SMELL_ORDER_CLEAN_THEN_ISO,
    SMELL_ORDER_ISO_THEN_CLEAN,
]
ASSIGNABLE_TASK_ORDER_GROUPS = [
    TASK_ORDER_IAT_BLOCK_FIRST,
    TASK_ORDER_AAT_BLOCK_FIRST,
]
ASSIGNABLE_IAT_VERSIONS = [1, 2, 3, 4]
ASSIGNABLE_GROUP_COMBINATIONS = [
    (smell_order_group, task_order_group)
    for smell_order_group in ASSIGNABLE_SMELL_ORDER_GROUPS
    for task_order_group in ASSIGNABLE_TASK_ORDER_GROUPS
]


@dataclass(frozen=True)
class ParticipantLogData:
    participant_id: str
    participant_display_id: str = ""
    participant_name_or_alias: str = ""
    phone_last4: str = ""
    session_number: int | None = None
    session_date: str | date | datetime | None = None
    session_condition: str = ""
    session_researcher: str = ""
    notes: str = ""
    smell_order_group: str = ""
    task_order_group: str = ""
    counterbalance_code: str = ""


@dataclass(frozen=True)
class RunLogData:
    participant_id: str
    participant_alias: str = ""
    task: str = ""
    script_type: str = ""
    task_order_group: str = ""
    run_timestamp: str | datetime | None = None
    session_number: int | None = None
    researcher: str = ""
    condition: str = ""
    iat_version: int | str | None = None
    iat_order_condition: str = ""
    iat_side_condition: str = ""
    aat_counterbalance: str = ""
    output_filename: str = ""
    output_path: str = ""
    status: str = "completed"
    notes: str = ""
    smell_order_group: str = ""
    counterbalance_code: str = ""


@dataclass(frozen=True)
class ParticipantAssignment:
    smell_order_group: str
    task_order_group: str
    source: str


@dataclass(frozen=True)
class LogbookWriteResult:
    logbook_path: Path
    participant_row: int
    participant_created: bool
    run_row: int | None
    warnings: tuple[str, ...]


def _load_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "openpyxl is required for workbook automation. Install it in the experiment environment "
            "before using the logbook helper."
        ) from exc
    return Workbook, load_workbook


def _normalize_value(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def validate_run_status(status: str) -> str:
    normalized = _normalize_value(status).strip().lower()
    if normalized not in VALID_RUN_STATUSES:
        raise ValueError(
            f"Unsupported run status: {status!r}. Expected one of {sorted(VALID_RUN_STATUSES)}."
        )
    return normalized


def _sheet_header_values(worksheet) -> list[str]:
    if worksheet.max_row == 1 and all(cell.value is None for cell in worksheet[1]):
        return []
    return [_normalize_value(cell.value) for cell in worksheet[1]]


def _ensure_sheet_headers(worksheet, headers: list[str]) -> None:
    current_headers = _sheet_header_values(worksheet)
    if not current_headers:
        for column_index, header in enumerate(headers, start=1):
            worksheet.cell(row=1, column=column_index, value=header)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        return
    if current_headers in (LEGACY_PARTICIPANTS_HEADERS, LEGACY_RUNS_HEADERS):
        for column_index, header in enumerate(headers[len(current_headers):], start=len(current_headers) + 1):
            worksheet.cell(row=1, column=column_index, value=header)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        return
    if current_headers != headers:
        raise RuntimeError(
            f"Sheet '{worksheet.title}' has unexpected headers.\n"
            f"Expected: {headers}\n"
            f"Found: {current_headers}"
        )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _clear_sheet_data(worksheet) -> None:
    if worksheet.max_row >= 2:
        worksheet.delete_rows(2, worksheet.max_row - 1)


def _iat_version_fields(iat_version: int) -> tuple[str, str]:
    if iat_version not in ASSIGNABLE_IAT_VERSIONS:
        raise ValueError(f"Unsupported IAT version: {iat_version}.")
    order_condition = "disease_first" if iat_version in (1, 2) else "danger_first"
    side_condition = "combined_left" if iat_version in (1, 3) else "combined_right"
    return order_condition, side_condition


def _aat_block_order(aat_counterbalance: str) -> str:
    return "congruent_first" if aat_counterbalance == AAT_COUNTERBALANCE_A else "incongruent_first"


def _aat_counterbalance_for_iat_version(iat_version: int) -> str:
    if iat_version not in ASSIGNABLE_IAT_VERSIONS:
        raise ValueError(f"Unsupported IAT version: {iat_version}.")
    return AAT_COUNTERBALANCE_A if iat_version in (2, 4) else assign_aat_counterbalance(iat_version)


def _counterbalance_code(task_order_group: str, smell_order_group: str, iat_version: int) -> str:
    validate_task_order_group(task_order_group)
    validate_smell_order_group(smell_order_group)
    if iat_version not in ASSIGNABLE_IAT_VERSIONS:
        raise ValueError(f"Unsupported IAT version: {iat_version}.")
    task_index = ASSIGNABLE_TASK_ORDER_GROUPS.index(task_order_group)
    smell_index = ASSIGNABLE_SMELL_ORDER_GROUPS.index(smell_order_group)
    code_number = task_index * (len(ASSIGNABLE_SMELL_ORDER_GROUPS) * len(ASSIGNABLE_IAT_VERSIONS))
    code_number += smell_index * len(ASSIGNABLE_IAT_VERSIONS)
    code_number += iat_version
    return f"CB{code_number:02d}"


def _participant_counterbalance_bundle(
    participant_id: str,
    smell_order_group: str,
    task_order_group: str,
) -> dict[str, object]:
    participant = parse_participant_last4(participant_id)
    iat_version = assign_iat_version(participant.numeric_id)
    iat_order_condition, iat_side_condition = _iat_version_fields(iat_version)
    aat_counterbalance = assign_aat_counterbalance(participant.numeric_id)
    odor_1, odor_2, odor_3 = smell_condition_order(smell_order_group)
    return {
        "counterbalance_code": _counterbalance_code(task_order_group, smell_order_group, iat_version),
        "iat_version": iat_version,
        "iat_order_condition": iat_order_condition,
        "iat_side_condition": iat_side_condition,
        "aat_counterbalance": aat_counterbalance,
        "aat_block_order": _aat_block_order(aat_counterbalance),
        "odor_1": odor_1,
        "odor_2": odor_2,
        "odor_3": odor_3,
    }


def _resolve_counterbalance_bundle(
    participant_id: str,
    smell_order_group: str,
    task_order_group: str,
) -> dict[str, object]:
    participant_id = _normalize_value(participant_id).strip()
    smell_order_group = _safe_group_value(
        smell_order_group,
        validator=validate_smell_order_group,
    )
    task_order_group = _safe_group_value(
        task_order_group,
        validator=validate_task_order_group,
    )
    if not participant_id or not smell_order_group or not task_order_group:
        return {}
    try:
        return _participant_counterbalance_bundle(participant_id, smell_order_group, task_order_group)
    except ValueError:
        return {}


def _counterbalancing_key_rows() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for task_order_group in ASSIGNABLE_TASK_ORDER_GROUPS:
        for smell_order_group in ASSIGNABLE_SMELL_ORDER_GROUPS:
            odor_1, odor_2, odor_3 = smell_condition_order(smell_order_group)
            for iat_version in ASSIGNABLE_IAT_VERSIONS:
                iat_order_condition, iat_side_condition = _iat_version_fields(iat_version)
                aat_counterbalance = _aat_counterbalance_for_iat_version(iat_version)
                rows.append(
                    {
                        "counterbalance_code": _counterbalance_code(
                            task_order_group,
                            smell_order_group,
                            iat_version,
                        ),
                        "task_order_group": task_order_group,
                        "smell_order_group": smell_order_group,
                        "odor_1": odor_1,
                        "odor_2": odor_2,
                        "odor_3": odor_3,
                        "iat_version": iat_version,
                        "iat_order_condition": iat_order_condition,
                        "iat_side_condition": iat_side_condition,
                        "aat_counterbalance": aat_counterbalance,
                        "aat_block_order": _aat_block_order(aat_counterbalance),
                    }
                )
    return rows


def _write_sheet_rows(worksheet, headers: list[str], rows: list[dict[str, object]]) -> None:
    _ensure_sheet_headers(worksheet, headers)
    _clear_sheet_data(worksheet)
    header_map = _header_map(headers)
    for row_index, row in enumerate(rows, start=2):
        for header in headers:
            worksheet.cell(
                row=row_index,
                column=header_map[header],
                value=_normalize_value(row.get(header, "")),
            )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _backfill_counterbalance_codes(workbook) -> None:
    participants_ws = workbook[SHEET_PARTICIPANTS]
    participant_header_map = _sheet_header_map(participants_ws, PARTICIPANTS_HEADERS)
    for row_index in range(2, participants_ws.max_row + 1):
        row = _worksheet_row_as_dict(participants_ws, row_index, PARTICIPANTS_HEADERS)
        if not any(str(value).strip() for value in row.values()):
            continue
        bundle = _resolve_counterbalance_bundle(
            row.get("participant_id", ""),
            row.get("smell_order_group", ""),
            row.get("task_order_group", ""),
        )
        if bundle:
            participants_ws.cell(
                row=row_index,
                column=participant_header_map["counterbalance_code"],
                value=bundle["counterbalance_code"],
            )

    runs_ws = workbook[SHEET_RUNS]
    run_header_map = _sheet_header_map(runs_ws, RUNS_HEADERS)
    for row_index in range(2, runs_ws.max_row + 1):
        row = _worksheet_row_as_dict(runs_ws, row_index, RUNS_HEADERS)
        if not any(str(value).strip() for value in row.values()):
            continue
        bundle = _resolve_counterbalance_bundle(
            row.get("participant_id", ""),
            row.get("smell_order_group", ""),
            row.get("task_order_group", ""),
        )
        if bundle:
            runs_ws.cell(
                row=row_index,
                column=run_header_map["counterbalance_code"],
                value=bundle["counterbalance_code"],
            )


def _participant_plan_rows(workbook) -> list[dict[str, object]]:
    participants_ws = workbook[SHEET_PARTICIPANTS]
    runs_by_participant: dict[str, list[dict[str, object]]] = {}
    for run_row in _load_run_rows(workbook[SHEET_RUNS]):
        participant_id = _normalize_value(run_row.get("participant_id", "")).strip()
        if participant_id:
            runs_by_participant.setdefault(participant_id, []).append(run_row)

    rows: list[dict[str, object]] = []
    for row_index in range(2, participants_ws.max_row + 1):
        participant_row = _worksheet_row_as_dict(participants_ws, row_index, PARTICIPANTS_HEADERS)
        if not any(str(value).strip() for value in participant_row.values()):
            continue

        participant_id = _normalize_value(participant_row.get("participant_id", "")).strip()
        if not participant_id:
            continue

        smell_order_group = _safe_group_value(
            participant_row.get("smell_order_group", ""),
            validator=validate_smell_order_group,
        )
        task_order_group = _safe_group_value(
            participant_row.get("task_order_group", ""),
            validator=validate_task_order_group,
        )
        bundle = _resolve_counterbalance_bundle(participant_id, smell_order_group, task_order_group)

        participant_runs = runs_by_participant.get(participant_id, [])
        latest_status = ""
        participant_alias = _normalize_value(participant_row.get("participant_name_or_alias", "")).strip()
        if participant_runs:
            latest_status = _normalize_value(participant_runs[-1].get("status", "")).strip()
            if not participant_alias:
                participant_alias = _normalize_value(participant_runs[-1].get("participant_alias", "")).strip()

        rows.append(
            {
                "participant_id": participant_id,
                "participant_alias": participant_alias,
                "counterbalance_code": bundle.get(
                    "counterbalance_code",
                    _normalize_value(participant_row.get("counterbalance_code", "")).strip(),
                ),
                "smell_order_group": smell_order_group,
                "odor_1": bundle.get("odor_1", ""),
                "odor_2": bundle.get("odor_2", ""),
                "odor_3": bundle.get("odor_3", ""),
                "task_order_group": task_order_group,
                "iat_version": bundle.get("iat_version", ""),
                "iat_order_condition": bundle.get("iat_order_condition", ""),
                "iat_side_condition": bundle.get("iat_side_condition", ""),
                "aat_counterbalance": bundle.get("aat_counterbalance", ""),
                "aat_block_order": bundle.get("aat_block_order", ""),
                "run_count": len(participant_runs),
                "latest_status": latest_status,
                "notes": _normalize_value(participant_row.get("notes", "")).strip(),
            }
        )
    return rows


def _refresh_counterbalance_views(workbook) -> None:
    _backfill_counterbalance_codes(workbook)
    _write_sheet_rows(
        workbook[SHEET_COUNTERBALANCING_KEY],
        COUNTERBALANCING_KEY_HEADERS,
        _counterbalancing_key_rows(),
    )
    _write_sheet_rows(
        workbook[SHEET_PARTICIPANT_PLAN],
        PARTICIPANT_PLAN_HEADERS,
        _participant_plan_rows(workbook),
    )


def ensure_logbook_template(logbook_path: Path | None = None) -> Path:
    Workbook, load_workbook = _load_openpyxl()
    path = Path(logbook_path or get_logbook_path())
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        workbook.active.title = SHEET_PARTICIPANTS

    for sheet_name in REQUIRED_SHEETS:
        if sheet_name not in workbook.sheetnames:
            workbook.create_sheet(title=sheet_name)

    _ensure_sheet_headers(workbook[SHEET_PARTICIPANTS], PARTICIPANTS_HEADERS)
    _ensure_sheet_headers(workbook[SHEET_RUNS], RUNS_HEADERS)
    _ensure_sheet_headers(workbook[SHEET_COUNTERBALANCING_KEY], COUNTERBALANCING_KEY_HEADERS)
    _ensure_sheet_headers(workbook[SHEET_PARTICIPANT_PLAN], PARTICIPANT_PLAN_HEADERS)

    workbook[SHEET_IAT_BALANCE].freeze_panes = "A1"
    workbook[SHEET_AAT_BALANCE].freeze_panes = "A1"
    _refresh_counterbalance_views(workbook)

    workbook.save(path)
    workbook.close()
    return path


def _header_map(headers: list[str]) -> dict[str, int]:
    return {header: index for index, header in enumerate(headers, start=1)}


def _sheet_header_map(worksheet, expected_headers: list[str]) -> dict[str, int]:
    current_headers = _sheet_header_values(worksheet)
    if current_headers != expected_headers:
        raise RuntimeError(
            f"Sheet '{worksheet.title}' headers do not match the expected structure."
        )
    return _header_map(expected_headers)


def _worksheet_row_as_dict(worksheet, row_index: int, headers: list[str]) -> dict[str, str]:
    values = {}
    for column_index, header in enumerate(headers, start=1):
        values[header] = _normalize_value(worksheet.cell(row=row_index, column=column_index).value)
    return values


def _set_cell_if_value(worksheet, row_index: int, column_index: int, value) -> None:
    normalized = _normalize_value(value)
    if normalized:
        worksheet.cell(row=row_index, column=column_index, value=normalized)


def _clear_sheet(worksheet) -> None:
    if worksheet.max_row > 0:
        worksheet.delete_rows(1, worksheet.max_row)


def _load_run_rows(runs_ws) -> list[dict[str, object]]:
    _sheet_header_map(runs_ws, RUNS_HEADERS)
    rows: list[dict[str, object]] = []
    for row_index in range(2, runs_ws.max_row + 1):
        row = _worksheet_row_as_dict(runs_ws, row_index, RUNS_HEADERS)
        if not any(str(value).strip() for value in row.values()):
            continue
        rows.append(row)
    return rows


def _count_run_values(rows: list[dict[str, object]], field_name: str) -> Counter:
    counts: Counter = Counter()
    for row in rows:
        value = row.get(field_name, "")
        if value in ("", None):
            continue
        counts[value] += 1
    return counts


def _write_balance_rows(worksheet, title: str, rows: list[dict[str, object]]) -> None:
    _clear_sheet(worksheet)
    worksheet.cell(row=1, column=1, value=title)
    worksheet.cell(row=2, column=1, value="generated_at")
    worksheet.cell(row=2, column=2, value=datetime.now().isoformat(timespec="seconds"))
    for column_index, header in enumerate(BALANCE_HEADERS, start=1):
        worksheet.cell(row=4, column=column_index, value=header)

    for row_index, row in enumerate(rows, start=5):
        worksheet.cell(row=row_index, column=1, value=row["metric"])
        worksheet.cell(row=row_index, column=2, value=row["value"])
        worksheet.cell(row=row_index, column=3, value=row["count"])

    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = worksheet.dimensions


def _build_balance_rows(
    runs: list[dict[str, object]],
    *,
    metrics: list[str],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = [
        {"metric": "total_runs", "value": "all", "count": len(runs)}
    ]
    for metric in metrics:
        counts = _count_run_values(runs, metric)
        for value in sorted(counts, key=lambda item: str(item)):
            rows.append({"metric": metric, "value": value, "count": counts[value]})
    return rows


def _session_field_names(session_number: int) -> tuple[str, str, str]:
    if session_number not in (1, 2, 3):
        raise ValueError(f"session_number must be 1, 2, or 3 to map into Participants, got {session_number}.")
    return (
        f"session_{session_number}_researcher",
        f"session_{session_number}_date",
        f"session_{session_number}_condition",
    )


def _safe_group_value(value: str, *, validator) -> str:
    normalized = _normalize_value(value).strip()
    if not normalized:
        return ""
    try:
        validator(normalized)
    except ValueError:
        return ""
    return normalized


def _find_participant_rows(participants_ws, participant_id: str) -> list[int]:
    header_map = _sheet_header_map(participants_ws, PARTICIPANTS_HEADERS)
    participant_col = header_map["participant_id"]
    matches: list[int] = []
    for row_index in range(2, participants_ws.max_row + 1):
        cell_value = _normalize_value(participants_ws.cell(row=row_index, column=participant_col).value)
        if cell_value == participant_id:
            matches.append(row_index)
    return matches


def _participant_group_counts(participants_ws) -> Counter:
    counts: Counter = Counter()
    for row_index in range(2, participants_ws.max_row + 1):
        row = _worksheet_row_as_dict(participants_ws, row_index, PARTICIPANTS_HEADERS)
        if not any(str(value).strip() for value in row.values()):
            continue
        smell_order_group = _safe_group_value(
            row.get("smell_order_group", ""),
            validator=validate_smell_order_group,
        )
        task_order_group = _safe_group_value(
            row.get("task_order_group", ""),
            validator=validate_task_order_group,
        )
        if smell_order_group and task_order_group:
            counts[(smell_order_group, task_order_group)] += 1
    return counts


def resolve_participant_assignment(
    participant_id: str,
    *,
    logbook_path: Path | None = None,
) -> ParticipantAssignment:
    _, load_workbook = _load_openpyxl()
    path = ensure_logbook_template(logbook_path)
    workbook = load_workbook(path)
    participants_ws = workbook[SHEET_PARTICIPANTS]

    matched_rows = _find_participant_rows(participants_ws, _normalize_value(participant_id).strip())
    existing_smell_order_group = ""
    existing_task_order_group = ""
    if matched_rows:
        row = _worksheet_row_as_dict(participants_ws, matched_rows[0], PARTICIPANTS_HEADERS)
        existing_smell_order_group = _safe_group_value(
            row.get("smell_order_group", ""),
            validator=validate_smell_order_group,
        )
        existing_task_order_group = _safe_group_value(
            row.get("task_order_group", ""),
            validator=validate_task_order_group,
        )
        if existing_smell_order_group and existing_task_order_group:
            workbook.close()
            return ParticipantAssignment(
                smell_order_group=existing_smell_order_group,
                task_order_group=existing_task_order_group,
                source="existing_participant_row",
            )

    group_counts = _participant_group_counts(participants_ws)
    workbook.close()

    candidate_combinations = [
        combination
        for combination in ASSIGNABLE_GROUP_COMBINATIONS
        if (not existing_smell_order_group or combination[0] == existing_smell_order_group)
        and (not existing_task_order_group or combination[1] == existing_task_order_group)
    ]
    if not candidate_combinations:
        candidate_combinations = ASSIGNABLE_GROUP_COMBINATIONS

    selected_smell_order_group, selected_task_order_group = min(
        candidate_combinations,
        key=lambda combination: (group_counts[combination], ASSIGNABLE_GROUP_COMBINATIONS.index(combination)),
    )
    return ParticipantAssignment(
        smell_order_group=selected_smell_order_group,
        task_order_group=selected_task_order_group,
        source="auto_balanced_assignment",
    )


def _next_run_id(runs_ws) -> int:
    if runs_ws.max_row <= 1:
        return 1
    existing_run_ids = []
    for row_index in range(2, runs_ws.max_row + 1):
        value = runs_ws.cell(row=row_index, column=1).value
        try:
            existing_run_ids.append(int(value))
        except (TypeError, ValueError):
            continue
    return (max(existing_run_ids) + 1) if existing_run_ids else 1


def ensure_participant_row(
    participant: ParticipantLogData,
    *,
    logbook_path: Path | None = None,
) -> LogbookWriteResult:
    _, load_workbook = _load_openpyxl()
    path = ensure_logbook_template(logbook_path)
    workbook = load_workbook(path)
    participants_ws = workbook[SHEET_PARTICIPANTS]
    header_map = _sheet_header_map(participants_ws, PARTICIPANTS_HEADERS)

    participant_id = _normalize_value(participant.participant_id).strip()
    if not participant_id:
        raise ValueError("participant_id is required.")

    participant_display_id = _normalize_value(participant.participant_display_id).strip() or participant_id
    participant_alias = _normalize_value(participant.participant_name_or_alias).strip()
    phone_last4 = _normalize_value(participant.phone_last4).strip() or participant_id

    matched_rows = _find_participant_rows(participants_ws, participant_id)
    warning_messages: list[str] = []
    if len(matched_rows) > 1:
        warning_message = (
            f"Duplicate participant_id '{participant_id}' detected in Participants sheet at rows "
            f"{', '.join(str(row) for row in matched_rows)}."
        )
        warnings.warn(warning_message)
        warning_messages.append(warning_message)

    if matched_rows:
        row_index = matched_rows[0]
        participant_created = False
    else:
        row_index = participants_ws.max_row + 1
        participants_ws.cell(row=row_index, column=header_map["participant_id"], value=participant_id)
        participant_created = True

    existing_row = (
        _worksheet_row_as_dict(participants_ws, row_index, PARTICIPANTS_HEADERS)
        if matched_rows
        else {}
    )

    _set_cell_if_value(participants_ws, row_index, header_map["participant_display_id"], participant_display_id)
    _set_cell_if_value(participants_ws, row_index, header_map["participant_name_or_alias"], participant_alias)
    _set_cell_if_value(participants_ws, row_index, header_map["phone_last4"], phone_last4)
    _set_cell_if_value(participants_ws, row_index, header_map["notes"], participant.notes)
    smell_order_group = _normalize_value(participant.smell_order_group).strip()
    if smell_order_group:
        validate_smell_order_group(smell_order_group)
    task_order_group = _normalize_value(participant.task_order_group).strip()
    if task_order_group:
        validate_task_order_group(task_order_group)
    if not smell_order_group:
        smell_order_group = _safe_group_value(
            existing_row.get("smell_order_group", ""),
            validator=validate_smell_order_group,
        )
    if not task_order_group:
        task_order_group = _safe_group_value(
            existing_row.get("task_order_group", ""),
            validator=validate_task_order_group,
        )
    _set_cell_if_value(participants_ws, row_index, header_map["smell_order_group"], smell_order_group)
    _set_cell_if_value(participants_ws, row_index, header_map["task_order_group"], task_order_group)
    counterbalance_code = _normalize_value(participant.counterbalance_code).strip()
    bundle = _resolve_counterbalance_bundle(participant_id, smell_order_group, task_order_group)
    if bundle:
        derived_counterbalance_code = _normalize_value(bundle["counterbalance_code"]).strip()
        if counterbalance_code and counterbalance_code != derived_counterbalance_code:
            raise ValueError(
                "counterbalance_code does not match the participant assignment implied by "
                f"{participant_id=}, {smell_order_group=}, and {task_order_group=}."
            )
        counterbalance_code = derived_counterbalance_code
    elif not counterbalance_code:
        counterbalance_code = _normalize_value(existing_row.get("counterbalance_code", "")).strip()
    _set_cell_if_value(participants_ws, row_index, header_map["counterbalance_code"], counterbalance_code)

    if participant.session_number in (1, 2, 3):
        session_researcher_field, session_date_field, session_condition_field = _session_field_names(
            participant.session_number
        )
        if participant.session_condition:
            validate_condition(participant.session_condition)
        _set_cell_if_value(
            participants_ws,
            row_index,
            header_map[session_researcher_field],
            participant.session_researcher,
        )
        _set_cell_if_value(
            participants_ws,
            row_index,
            header_map[session_date_field],
            participant.session_date,
        )
        _set_cell_if_value(
            participants_ws,
            row_index,
            header_map[session_condition_field],
            participant.session_condition,
        )

    participants_ws.auto_filter.ref = participants_ws.dimensions
    workbook.save(path)
    workbook.close()
    return LogbookWriteResult(
        logbook_path=path,
        participant_row=row_index,
        participant_created=participant_created,
        run_row=None,
        warnings=tuple(warning_messages),
    )


def append_run_row(
    run: RunLogData,
    *,
    logbook_path: Path | None = None,
) -> LogbookWriteResult:
    _, load_workbook = _load_openpyxl()
    path = ensure_logbook_template(logbook_path)
    workbook = load_workbook(path)
    runs_ws = workbook[SHEET_RUNS]
    header_map = _sheet_header_map(runs_ws, RUNS_HEADERS)

    participant_id = _normalize_value(run.participant_id).strip()
    if not participant_id:
        raise ValueError("participant_id is required for Runs entries.")
    validate_task(run.task)
    validate_script_type(run.script_type)
    validate_task_order_group(run.task_order_group)
    validate_run_status(run.status)
    if run.smell_order_group:
        validate_smell_order_group(run.smell_order_group)
    if run.condition:
        validate_condition(run.condition)

    bundle = _resolve_counterbalance_bundle(
        participant_id,
        _normalize_value(run.smell_order_group).strip(),
        run.task_order_group,
    )
    counterbalance_code = _normalize_value(run.counterbalance_code).strip()
    if bundle:
        derived_counterbalance_code = _normalize_value(bundle["counterbalance_code"]).strip()
        if counterbalance_code and counterbalance_code != derived_counterbalance_code:
            raise ValueError(
                "counterbalance_code does not match the run assignment implied by "
                f"{participant_id=}, smell_order_group={run.smell_order_group!r}, and "
                f"task_order_group={run.task_order_group!r}."
            )
        counterbalance_code = derived_counterbalance_code

    row_index = runs_ws.max_row + 1
    runs_ws.cell(row=row_index, column=header_map["run_id"], value=_next_run_id(runs_ws))
    row_values = {
        "participant_id": participant_id,
        "participant_alias": _normalize_value(run.participant_alias).strip(),
        "task": run.task,
        "script_type": run.script_type,
        "task_order_group": run.task_order_group,
        "run_timestamp": _normalize_value(run.run_timestamp) or datetime.now().isoformat(timespec="seconds"),
        "session_number": _normalize_value(run.session_number),
        "researcher": _normalize_value(run.researcher).strip(),
        "condition": _normalize_value(run.condition).strip(),
        "iat_version": _normalize_value(run.iat_version),
        "iat_order_condition": _normalize_value(run.iat_order_condition).strip(),
        "iat_side_condition": _normalize_value(run.iat_side_condition).strip(),
        "aat_counterbalance": _normalize_value(run.aat_counterbalance).strip(),
        "output_filename": _normalize_value(run.output_filename).strip(),
        "output_path": _normalize_value(run.output_path).strip(),
        "status": validate_run_status(run.status),
        "notes": _normalize_value(run.notes).strip(),
        "smell_order_group": _normalize_value(run.smell_order_group).strip(),
        "counterbalance_code": counterbalance_code,
    }

    for header, value in row_values.items():
        runs_ws.cell(row=row_index, column=header_map[header], value=value)

    runs_ws.auto_filter.ref = runs_ws.dimensions
    workbook.save(path)
    workbook.close()
    return LogbookWriteResult(
        logbook_path=path,
        participant_row=-1,
        participant_created=False,
        run_row=row_index,
        warnings=(),
    )


def record_run(
    participant: ParticipantLogData,
    run: RunLogData,
    *,
    logbook_path: Path | None = None,
) -> LogbookWriteResult:
    participant_result = ensure_participant_row(participant, logbook_path=logbook_path)
    run_result = append_run_row(run, logbook_path=participant_result.logbook_path)
    warning_messages = list(participant_result.warnings + run_result.warnings)
    try:
        refresh_balance_summaries(run_result.logbook_path)
    except Exception as exc:
        warning_message = f"Balance summary refresh failed: {type(exc).__name__}: {exc}"
        warnings.warn(warning_message)
        warning_messages.append(warning_message)
    return LogbookWriteResult(
        logbook_path=participant_result.logbook_path,
        participant_row=participant_result.participant_row,
        participant_created=participant_result.participant_created,
        run_row=run_result.run_row,
        warnings=tuple(warning_messages),
    )


def refresh_balance_summaries(logbook_path: Path | None = None) -> Path:
    _, load_workbook = _load_openpyxl()
    path = ensure_logbook_template(logbook_path)
    workbook = load_workbook(path)
    runs_ws = workbook[SHEET_RUNS]
    run_rows = _load_run_rows(runs_ws)

    iat_rows = [row for row in run_rows if row["task"] == "IAT"]
    aat_rows = [row for row in run_rows if row["task"] == "AAT"]

    iat_balance_rows = _build_balance_rows(
        iat_rows,
        metrics=[
            "iat_version",
            "iat_order_condition",
            "iat_side_condition",
            "condition",
            "smell_order_group",
            "task_order_group",
        ],
    )
    aat_balance_rows = _build_balance_rows(
        aat_rows,
        metrics=[
            "aat_counterbalance",
            "condition",
            "script_type",
            "smell_order_group",
            "task_order_group",
        ],
    )

    _write_balance_rows(workbook[SHEET_IAT_BALANCE], "IAT Balance Summary", iat_balance_rows)
    _write_balance_rows(workbook[SHEET_AAT_BALANCE], "AAT Balance Summary", aat_balance_rows)
    _refresh_counterbalance_views(workbook)

    workbook.save(path)
    workbook.close()
    return path
