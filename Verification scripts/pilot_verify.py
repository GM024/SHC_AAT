from __future__ import annotations

import argparse
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path


VERIFICATION_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = VERIFICATION_DIR.parent
CACHE_DIR = PROJECT_ROOT / "cache"
CACHE_DIR.mkdir(parents=True, exist_ok=True)
os.environ["PYTHONPYCACHEPREFIX"] = str(CACHE_DIR)
sys.pycache_prefix = str(CACHE_DIR)
IAT_DIR = PROJECT_ROOT / "Implicit Association Task" / "IAT Task Script"
AAT_DIR = PROJECT_ROOT / "Approach Avoidance Task" / "AAT Task Script"
for path in (VERIFICATION_DIR, PROJECT_ROOT, IAT_DIR, AAT_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from openpyxl import load_workbook

import aat_core
import iat_core
import iat_validate
from shared.experiment_meta import (
    CONDITION_CLEAN_ODOUR,
    CONDITION_CONTROL_NO_SMELL,
    CONDITION_ISOVALERIC_ACID,
    SCRIPT_TYPE_MAIN,
    SCRIPT_TYPE_PRACTICE,
    SMELL_ORDER_CLEAN_THEN_ISO,
    TASK_ORDER_IAT_BLOCK_FIRST,
)
from shared.logbook import ensure_logbook_template, refresh_balance_summaries


def _summary_rows_to_map(worksheet) -> dict[str, Counter]:
    summary: dict[str, Counter] = defaultdict(Counter)
    for row_index in range(5, worksheet.max_row + 1):
        metric = worksheet.cell(row=row_index, column=1).value
        value = worksheet.cell(row=row_index, column=2).value
        count = worksheet.cell(row=row_index, column=3).value
        if metric in (None, ""):
            continue
        summary[str(metric)][value] = int(count)
    return summary


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def _build_pilot_logbook_path(base_dir: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return base_dir / f"MRP_Logbook_pilot_{timestamp}.xlsx"


def run_pilot(logbook_path: Path) -> Path:
    ensure_logbook_template(logbook_path)

    validation_lines = iat_validate.run_validation()
    _assert(validation_lines[-1] == "Validation passed.", "IAT dry-run validator did not pass.")

    iat_practice_setup = iat_core.build_run_setup(
        participant_last4="0123",
        script_type=SCRIPT_TYPE_PRACTICE,
        session_number=None,
        smell_order_group=SMELL_ORDER_CLEAN_THEN_ISO,
        task_order_group=TASK_ORDER_IAT_BLOCK_FIRST,
        participant_alias="P0123",
        researcher="GM",
    )
    iat_main_setup = iat_core.build_run_setup(
        participant_last4="0123",
        script_type=SCRIPT_TYPE_MAIN,
        session_number=2,
        smell_order_group=SMELL_ORDER_CLEAN_THEN_ISO,
        task_order_group=TASK_ORDER_IAT_BLOCK_FIRST,
        participant_alias="P0123",
        researcher="GM",
    )
    aat_practice_setup = aat_core.build_run_setup(
        participant_last4="0123",
        script_type=SCRIPT_TYPE_PRACTICE,
        session_number=None,
        smell_order_group=SMELL_ORDER_CLEAN_THEN_ISO,
        task_order_group=TASK_ORDER_IAT_BLOCK_FIRST,
        participant_alias="P0123",
        researcher="GM",
    )
    aat_main_setup = aat_core.build_run_setup(
        participant_last4="0123",
        script_type=SCRIPT_TYPE_MAIN,
        session_number=3,
        smell_order_group=SMELL_ORDER_CLEAN_THEN_ISO,
        task_order_group=TASK_ORDER_IAT_BLOCK_FIRST,
        participant_alias="P0123",
        researcher="GM",
    )

    _assert(iat_main_setup.run.condition == CONDITION_CLEAN_ODOUR, "IAT session 2 should resolve to clean_odour.")
    _assert(
        aat_main_setup.run.condition == CONDITION_ISOVALERIC_ACID,
        "AAT session 3 should resolve to isovaleric_acid for the clean-then-iso smell order.",
    )

    iat_practice_output = PROJECT_ROOT / "Implicit Association Task" / "data" / iat_core.build_output_filename(
        iat_practice_setup.run,
        timestamp=datetime(2026, 3, 22, 20, 0, 0),
        iat_version=iat_practice_setup.iat_version,
    )
    iat_main_output = PROJECT_ROOT / "Implicit Association Task" / "data" / iat_core.build_output_filename(
        iat_main_setup.run,
        timestamp=datetime(2026, 3, 22, 21, 0, 0),
        iat_version=iat_main_setup.iat_version,
    )
    aat_practice_output = PROJECT_ROOT / "Approach Avoidance Task" / "data" / aat_core.build_output_filename(
        aat_practice_setup.run,
        timestamp=datetime(2026, 3, 22, 22, 0, 0),
    )
    aat_main_output = PROJECT_ROOT / "Approach Avoidance Task" / "data" / aat_core.build_output_filename(
        aat_main_setup.run,
        timestamp=datetime(2026, 3, 22, 23, 0, 0),
    )

    iat_core.write_iat_logbook_entry(
        iat_practice_setup,
        out_csv=iat_practice_output,
        status="completed",
        run_timestamp=datetime(2026, 3, 22, 20, 0, 0),
        logbook_path=logbook_path,
    )
    iat_core.write_iat_logbook_entry(
        iat_main_setup,
        out_csv=iat_main_output,
        status="completed",
        run_timestamp=datetime(2026, 3, 22, 21, 0, 0),
        logbook_path=logbook_path,
    )
    aat_core.write_aat_logbook_entry(
        aat_practice_setup,
        output_path=aat_practice_output,
        status="completed",
        run_timestamp=datetime(2026, 3, 22, 22, 0, 0),
        logbook_path=logbook_path,
    )
    aat_core.write_aat_logbook_entry(
        aat_main_setup,
        output_path=aat_main_output,
        status="completed",
        run_timestamp=datetime(2026, 3, 22, 23, 0, 0),
        logbook_path=logbook_path,
    )

    refresh_balance_summaries(logbook_path)

    workbook = load_workbook(logbook_path)
    participants_ws = workbook["Participants"]
    runs_ws = workbook["Runs"]
    iat_balance_ws = workbook["IAT Balance"]
    aat_balance_ws = workbook["AAT Balance"]

    _assert(participants_ws.max_row == 2, "Pilot workbook should contain exactly one participant row.")
    _assert(runs_ws.max_row == 5, "Pilot workbook should contain exactly four run rows.")
    _assert(participants_ws["A2"].value == "0123", "Participant ID was not written correctly.")
    _assert(participants_ws["C2"].value == "P0123", "Participant alias was not written correctly.")
    _assert(participants_ws["E2"].value == "GM", "Session 2 researcher traceability is missing.")
    _assert(participants_ws["F2"].value == "GM", "Session 3 researcher traceability is missing.")
    _assert(participants_ws["L2"].value == CONDITION_CLEAN_ODOUR, "Session 2 condition traceability is incorrect.")
    _assert(
        participants_ws["M2"].value == CONDITION_ISOVALERIC_ACID,
        "Session 3 condition traceability is incorrect.",
    )

    run_rows = []
    for row_index in range(2, runs_ws.max_row + 1):
        run_rows.append(
            {
                "task": runs_ws.cell(row=row_index, column=4).value,
                "script_type": runs_ws.cell(row=row_index, column=5).value,
                "task_order_group": runs_ws.cell(row=row_index, column=6).value,
                "condition": runs_ws.cell(row=row_index, column=10).value,
                "output_filename": runs_ws.cell(row=row_index, column=15).value,
                "status": runs_ws.cell(row=row_index, column=17).value,
            }
        )

    _assert({row["task"] for row in run_rows} == {"IAT", "AAT"}, "Runs sheet mixes task labels incorrectly.")
    _assert(
        Counter(row["script_type"] for row in run_rows) == Counter({"practice": 2, "main": 2}),
        "Practice/main run counts are incorrect.",
    )
    _assert(
        all(row["task_order_group"] == TASK_ORDER_IAT_BLOCK_FIRST for row in run_rows),
        "Task-order group was not preserved across pilot runs.",
    )
    _assert(
        all(row["status"] == "completed" for row in run_rows),
        "Pilot verification should have completed-only rows.",
    )
    _assert(
        all(
            ("practice" in row["output_filename"]) == (row["script_type"] == "practice")
            for row in run_rows
        ),
        "Practice/main output filenames do not match their script types.",
    )
    _assert(
        all(
            row["output_filename"].startswith("iat_") if row["task"] == "IAT" else row["output_filename"].startswith("aat_")
            for row in run_rows
        ),
        "Task output filenames do not match their task labels.",
    )

    iat_summary = _summary_rows_to_map(iat_balance_ws)
    aat_summary = _summary_rows_to_map(aat_balance_ws)
    _assert(iat_summary["total_runs"]["all"] == 2, "IAT balance summary total is incorrect.")
    _assert(iat_summary["iat_version"][iat_main_setup.iat_version] == 2, "IAT version summary is incorrect.")
    _assert(iat_summary["condition"][CONDITION_CONTROL_NO_SMELL] == 1, "IAT control count is incorrect.")
    _assert(iat_summary["condition"][CONDITION_CLEAN_ODOUR] == 1, "IAT clean odour count is incorrect.")
    _assert(aat_summary["total_runs"]["all"] == 2, "AAT balance summary total is incorrect.")
    _assert(
        aat_summary["aat_counterbalance"][aat_main_setup.counterbalance_condition] == 2,
        "AAT counterbalance summary is incorrect.",
    )
    _assert(aat_summary["script_type"]["practice"] == 1, "AAT practice count is incorrect.")
    _assert(aat_summary["script_type"]["main"] == 1, "AAT main count is incorrect.")
    _assert(
        aat_summary["condition"][CONDITION_ISOVALERIC_ACID] == 1,
        "AAT isovaleric condition count is incorrect.",
    )

    workbook.close()
    return logbook_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Create a pilot workbook and verify the full MRP logbook flow.")
    parser.add_argument(
        "--logbook",
        type=Path,
        default=None,
        help="Optional workbook path. Defaults to logbook/pilot/MRP_Logbook_pilot_<timestamp>.xlsx.",
    )
    args = parser.parse_args()

    if args.logbook is None:
        pilot_dir = PROJECT_ROOT.parent / "logbook" / "pilot"
        pilot_dir.mkdir(parents=True, exist_ok=True)
        logbook_path = _build_pilot_logbook_path(pilot_dir)
    else:
        logbook_path = args.logbook

    run_pilot(logbook_path)
    print(logbook_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
